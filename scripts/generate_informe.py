#!/usr/bin/env python3
"""genera informe ejecutivo en formato docx."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "entregables" / "informe_evaluacion.docx"
CAPTURAS = ROOT / "docs" / "capturas"

REPO_URL = "https://github.com/javier-tc/Evaluacion3_ProgramacionCD"
FECHA = date.today().strftime("%d de %B de %Y").replace(
    "January", "enero").replace("February", "febrero").replace(
    "March", "marzo").replace("April", "abril").replace(
    "May", "mayo").replace("June", "junio").replace(
    "July", "julio").replace("August", "agosto").replace(
    "September", "septiembre").replace("October", "octubre").replace(
    "November", "noviembre").replace("December", "diciembre")


def _heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def _para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    return p


def _bullet(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    doc.add_paragraph()


def _image(doc, path, width=5.5):
    if Path(path).exists():
        doc.add_picture(str(path), width=Inches(width))
        doc.add_paragraph()


def build():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    #portada
    for _ in range(6):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(
        "Informe Ejecutivo\nSistema de Monitoreo y Análisis\nde Calidad del Aire en Santiago"
    )
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(0x1A, 0x47, 0x7A)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(f"\nEvaluación 3 — Programación y Ciencia de Datos\n{FECHA}").font.size = Pt(12)

    team = doc.add_paragraph()
    team.alignment = WD_ALIGN_PARAGRAPH.CENTER
    team.add_run(
        "\n\nIntegrantes: @javier-tc · @roz-ctrl\n"
        f"Repositorio: {REPO_URL}"
    )

    doc.add_page_break()

    #resumen ejecutivo
    _heading(doc, "Resumen ejecutivo")
    _para(doc, (
        "Este informe presenta una solución integral de ciencia de datos para el monitoreo "
        "y análisis de la calidad del aire en la Región Metropolitana. El sistema integra "
        "seis fuentes de información (cinco series de contaminantes y datos meteorológicos "
        "de Open-Meteo), las procesa mediante un pipeline ETL automatizado, las almacena en "
        "PostgreSQL y las expone a través de una API REST y dashboards interactivos orientados "
        "a distintas audiencias."
    ))
    _para(doc, (
        "Los resultados principales incluyen: un pipeline reproducible desplegable con un "
        "único comando Docker, validación de esquemas con Pydantic y Great Expectations, "
        "25 pruebas automatizadas con 72% de cobertura en módulos ETL y API, y cuatro vistas "
        "de dashboard diferenciadas (ejecutiva, analítica, operacional y autoridad ambiental)."
    ))
    _para(doc, (
        "Se recomienda adoptar esta arquitectura como base para reportes periódicos de "
        "calidad del aire, priorizando en una segunda fase la automatización programada "
        "del ETL y la integración de fuentes adicionales en tiempo real."
    ), bold=True)

    #contexto
    _heading(doc, "1. Contexto y objetivos de negocio")
    _para(doc, (
        "La contaminación atmosférica en Santiago constituye un desafío de salud pública y "
        "gestión ambiental. Las autoridades y organizaciones requieren información confiable, "
        "actualizada y presentada según el perfil del decisor: ejecutivos necesitan KPIs "
        "sintéticos, analistas requieren correlaciones y tendencias, operaciones necesita "
        "monitoreo del pipeline, y la autoridad ambiental necesita validar hipótesis científicas."
    ))
    _bullet(doc, [
        "Integrar al menos tres fuentes heterogéneas de datos ambientales y meteorológicos.",
        "Garantizar calidad y trazabilidad mediante validación de esquemas y logging.",
        "Exponer información mediante API desacoplada y dashboards por audiencia.",
        "Asegurar reproducibilidad y despliegue containerizado.",
    ])

    _heading(doc, "Hipótesis de investigación", level=2)
    _table(doc, ["Hipótesis", "Enunciado"], [
        ["H1", "Mayor temperatura produce mayor concentración de ozono (O3)."],
        ["H2", "La lluvia reduce las concentraciones de MP2.5 y MP10."],
        ["H3", "Mayor velocidad del viento reduce CO y NO2."],
        ["H4", "Existen patrones estacionales en los contaminantes medidos."],
    ])

    #solucion
    _heading(doc, "2. Solución implementada")
    _para(doc, (
        "La solución sigue una arquitectura por capas: fuentes de datos → ETL → PostgreSQL → "
        "FastAPI → Dash. Esta separación permite escalar componentes de forma independiente "
        "y garantiza que el dashboard no acceda directamente a la base de datos."
    ))
    _image(doc, CAPTURAS / "arquitectura.png")

    #fuentes
    _heading(doc, "3. Fuentes de datos y cobertura")
    _table(doc, ["#", "Fuente", "Tipo", "Variable"], [
        ["1", "co_ppm_ohiggins_last_year.csv", "CSV local", "CO (ppm)"],
        ["2", "mp10_ug-m3_ohiggins_last_year.csv", "CSV local", "MP10 (µg/m³)"],
        ["3", "mp25_ohiggins_ly.csv", "CSV local", "MP2.5 (µg/m³)"],
        ["4", "no2_ppb_ohiggins_ly.csv", "CSV local", "NO2 (ppb)"],
        ["5", "o3_ppb_ohiggins_ly.csv", "CSV local", "O3 (ppb)"],
        ["6", "Open-Meteo ERA5", "API REST", "Temperatura, lluvia, viento"],
    ])
    _para(doc, (
        "Período de análisis: 01-06-2025 al 01-06-2026. Limitación conocida: los CSV "
        "contienen datos diarios (hora 00:00) de la estación O'Higgins, referenciados "
        "como Santiago en el sistema."
    ))

    #etl
    _heading(doc, "4. Pipeline ETL")
    _image(doc, CAPTURAS / "pipeline_etl.png")
    _para(doc, (
        "El pipeline ejecuta cuatro etapas secuenciales con manejo de errores en cada capa:"
    ))
    _bullet(doc, [
        "Extract: lectura de 5 CSV con validación de encabezados y API Open-Meteo con reintentos (tenacity).",
        "Transform: coalesce de calidad, parseo de fechas, agregados diarios/mensuales y correlaciones Pearson.",
        "Validate: Pydantic (tipos, rangos) + Great Expectations (completitud). Modo estricto configurable.",
        "Load: inserción transaccional en PostgreSQL con rollback y registro en etl_execution_log.",
    ])
    _para(doc, (
        "Scripts: scripts/run_etl.py. Notebooks exploratorios: etl/notebooks/01_exploracion_fuentes.ipynb "
        "y 02_pipeline_etl_demo.ipynb."
    ))

    #dashboards
    _heading(doc, "5. Dashboards por audiencia")
    _table(doc, ["Vista", "Ruta", "Audiencia", "Valor de negocio"], [
        ["Ejecutiva", "/ejecutivo", "Directivos", "KPIs, ranking y evolución mensual"],
        ["Analítica", "/analitico", "Analistas", "Correlaciones, scatter, estacionalidad"],
        ["Operacional", "/operacional", "DevOps/SRE", "Estado API, ETL y calidad de datos"],
        ["Autoridad", "/presentacion", "Municipalidad", "Hipótesis e insights ambientales"],
    ])
    _image(doc, CAPTURAS / "dashboard_ejecutivo.png", 4.5)
    _image(doc, CAPTURAS / "dashboard_analitico.png", 4.5)

    #api
    _heading(doc, "6. API REST y exposición de datos")
    _para(doc, (
        "FastAPI expone 11 endpoints documentados en Swagger (/docs): mediciones de contaminación, "
        "clima, métricas agregadas, correlaciones, tendencias, calidad de datos y estado del ETL. "
        "La API actúa como capa de desacoplamiento entre almacenamiento y visualización."
    ))

    #docker
    _heading(doc, "7. Infraestructura Docker")
    _image(doc, CAPTURAS / "docker_compose.png")
    _para(doc, (
        "Docker Compose orquesta cuatro servicios con dependencias explícitas: PostgreSQL "
        "(healthcheck) → ETL (one-shot) → API (healthy) → Dashboard. Variables de entorno "
        "centralizadas en .env. Scripts de despliegue: docker/deploy.sh y docker/deploy.ps1."
    ))

    #calidad
    _heading(doc, "8. Calidad, testing y confiabilidad")
    _bullet(doc, [
        "25 pruebas automatizadas (pytest) — 100% exitosas.",
        "72% cobertura en módulos etl/ y api/.",
        "Logging estructurado con loguru (rotación 10 MB).",
        "Trazabilidad por run_id en etl_execution_log.",
        "Variable ETL_STRICT_VALIDATION para abortar carga ante fallos de validación.",
    ])

    #git
    _heading(doc, "9. Colaboración y gobernanza Git")
    _para(doc, (
        "El equipo adoptó Git Flow documentado en repo/GITFLOW.md con ramas main y develop, "
        "features con Pull Requests y convención de commits semánticos. Colaboradores: "
        "@javier-tc (ETL, API, dashboard, Docker) y @roz-ctrl (organización del repo, "
        "documentación del equipo). Evidencia en repo/COLABORACION.md y GitHub."
    ))

    #matriz
    _heading(doc, "10. Matriz de cumplimiento de la rúbrica")
    _table(doc, ["Requisito", "Estado", "Evidencia"], [
        ["Estructura de carpetas recomendada", "Cumple", "/etl, /dashboards, /docs, /api, /docker, /tests, /data, /repo"],
        ["Pipeline ETL ≥3 fuentes", "Cumple", "6 fuentes integradas"],
        ["Scripts y notebooks", "Cumple", "run_etl.py + 2 notebooks"],
        ["Validación de esquemas", "Cumple", "Pydantic + Great Expectations"],
        ["Manejo avanzado de errores", "Cumple", "tenacity, rollback, modo estricto"],
        ["Documentación técnica", "Cumple", "docs/ completo"],
        ["Dashboard por audiencia", "Cumple", "4 vistas Dash"],
        ["Git colaborativo", "Parcial", "Git Flow + PRs; historial en consolidación"],
        ["Docker + compose + env", "Cumple", "docker-compose.yml + .env.example"],
        ["Scripts de despliegue", "Cumple", "docker/deploy.sh, deploy.ps1"],
        ["Testing automatizado", "Cumple", "25 tests pytest, 72% cobertura"],
    ])

    #lecciones
    _heading(doc, "11. Lecciones aprendidas y mejoras futuras")
    _bullet(doc, [
        "La containerización con Docker Compose simplificó la demo end-to-end en evaluaciones.",
        "Separar API y dashboard facilitó el testing independiente de cada capa.",
        "La validación dual (Pydantic + GE) detectó inconsistencias tempranas en los CSV.",
        "Mejora propuesta: scheduler (Airflow/cron) para ETL periódico.",
        "Mejora propuesta: CI/CD con GitHub Actions para tests automáticos en cada PR.",
        "Mejora propuesta: tests de integración E2E con contenedores de prueba.",
    ])

    #anexos
    _heading(doc, "Anexos")
    _bullet(doc, [
        "API Swagger: http://localhost:8000/docs",
        "Dashboard: http://localhost:8050",
        "Inicio rápido: cp .env.example .env && docker compose up --build",
        "Documentación: docs/README.md",
        "Repositorio: " + REPO_URL,
    ])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Informe generado: {OUT}")


if __name__ == "__main__":
    build()
