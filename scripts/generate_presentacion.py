#!/usr/bin/env python3
"""genera presentacion de defensa en formato pptx."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "entregables" / "presentacion_evaluacion.pptx"
CAPTURAS = ROOT / "docs" / "capturas"

REPO_URL = "https://github.com/javier-tc/Evaluacion3_ProgramacionCD"
COLOR_TITLE = RGBColor(0x1A, 0x47, 0x7A)
COLOR_ACCENT = RGBColor(0x17, 0xA2, 0xB8)


def _title_slide(prs, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(8.4), Inches(1.5))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = COLOR_TITLE
    p.alignment = PP_ALIGN.CENTER
    if subtitle:
        box2 = slide.shapes.add_textbox(Inches(0.8), Inches(3.8), Inches(8.4), Inches(1.2))
        tf2 = box2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = subtitle
        p2.font.size = Pt(16)
        p2.alignment = PP_ALIGN.CENTER
    return slide


def _content_slide(prs, title, bullets, notes=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
    tp = title_box.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(28)
    tp.font.bold = True
    tp.font.color.rgb = COLOR_TITLE

    body = slide.shapes.add_textbox(Inches(0.7), Inches(1.2), Inches(8.6), Inches(5.5))
    tf = body.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.font.size = Pt(18)
        p.level = 0
        p.space_after = Pt(8)

    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def _image_slide(prs, title, image_path, bullets=None, notes=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.7))
    tp = title_box.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(26)
    tp.font.bold = True
    tp.font.color.rgb = COLOR_TITLE

    img_left = Inches(0.5)
    img_top = Inches(1.1)
    img_w = Inches(5.5) if bullets else Inches(8.5)
    if Path(image_path).exists():
        slide.shapes.add_picture(str(image_path), img_left, img_top, width=img_w)

    if bullets:
        body = slide.shapes.add_textbox(Inches(6.2), Inches(1.1), Inches(3.3), Inches(5.5))
        tf = body.text_frame
        for i, b in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = b
            p.font.size = Pt(14)

    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    _title_slide(
        prs,
        "Sistema de Monitoreo de Calidad del Aire",
        f"@javier-tc · @roz-ctrl\nEvaluación 3 — Programación y Ciencia de Datos\n{REPO_URL}",
    )

    _content_slide(prs, "Agenda", [
        "1. Contexto y objetivos",
        "2. Arquitectura y pipeline ETL",
        "3. Dashboards por audiencia",
        "4. Docker, Git y testing",
        "5. Lecciones aprendidas y demo en vivo",
    ])

    _content_slide(prs, "Problema y contexto", [
        "Santiago enfrenta desafíos de calidad del aire con impacto en salud pública.",
        "Múltiples audiencias requieren información diferenciada:",
        "  · Ejecutivos → KPIs y tendencias",
        "  · Analistas → correlaciones y patrones",
        "  · Operaciones → estado del pipeline",
        "  · Autoridad ambiental → validación de hipótesis",
    ], notes="Enfatizar el valor de negocio antes del detalle técnico.")

    _content_slide(prs, "Objetivos", [
        "Integrar ≥3 fuentes de datos ambientales y meteorológicos.",
        "Pipeline ETL robusto con validación y manejo de errores.",
        "API REST + dashboards interactivos Dash.",
        "Despliegue reproducible con Docker.",
        "Colaboración profesional con Git Flow.",
    ])

    _image_slide(
        prs, "Arquitectura técnica",
        CAPTURAS / "arquitectura.png",
        [
            "FastAPI: capa REST",
            "PostgreSQL: almacenamiento",
            "Dash: visualización",
            "Dashboard NO accede a BD",
        ],
        notes="Explicar decisiones: Python, SQLAlchemy, desacoplamiento por capas.",
    )

    _content_slide(prs, "Fuentes de datos (6)", [
        "5 CSV de contaminantes — estación O'Higgins:",
        "  CO, MP10, MP2.5, NO2, O3",
        "1 API Open-Meteo ERA5:",
        "  temperatura, lluvia, viento",
        "Período: jun 2025 — jun 2026",
    ])

    _image_slide(
        prs, "Pipeline ETL",
        CAPTURAS / "pipeline_etl.png",
        [
            "Extract: CSV + API",
            "Transform: agregados",
            "Validate: Pydantic + GE",
            "Load: PostgreSQL",
        ],
        notes="DEMO: docker compose up --build. Mostrar logs del ETL y etl_execution_log.",
    )

    _content_slide(prs, "Transformaciones y validación", [
        "Coalesce de calidad: validados > preliminares > no validados",
        "Agregados diarios y mensuales",
        "Correlaciones Pearson contaminante × clima",
        "Pydantic: tipos, rangos, coherencia temporal",
        "Great Expectations: completitud y no-negatividad",
        "ETL_STRICT_VALIDATION=true aborta carga si falla validación",
    ])

    _content_slide(prs, "Demo ETL — Guion", [
        "1. ./docker/deploy.ps1 (o docker compose up --build)",
        "2. Verificar postgres healthy: docker compose ps",
        "3. Observar logs ETL: docker compose logs etl",
        "4. Confirmar API: curl localhost:8000/health",
        "5. Abrir dashboard: localhost:8050",
        "6. Re-ejecutar ETL: ./docker/deploy.ps1 -EtlOnly",
    ], notes="Preferir demo en Docker. Tener terminal y navegador listos.")

    _content_slide(prs, "API REST (11 endpoints)", [
        "GET /pollution, /daily, /monthly",
        "GET /weather",
        "GET /analytics/correlation, /trends, /top-pollution-days",
        "GET /analytics/annual-averages, /etl-status, /data-quality",
        "Documentación interactiva: localhost:8000/docs",
    ])

    _image_slide(
        prs, "Dashboard — Vista ejecutiva",
        CAPTURAS / "dashboard_ejecutivo.png",
        ["5 KPI cards", "Ranking barras", "Evolución mensual"],
        notes="Adaptar discurso: ROI, alertas, decisiones rápidas.",
    )

    _image_slide(
        prs, "Dashboard — Vista analítica",
        CAPTURAS / "dashboard_analitico.png",
        ["Matriz correlación", "Scatter plots", "Estacionalidad"],
        notes="Audiencia técnica: profundizar en correlaciones y outliers.",
    )

    _image_slide(
        prs, "Dashboard — Vista operacional",
        CAPTURAS / "dashboard_operacional.png",
        ["Estado API/ETL", "Calidad datos", "Refresh 30s"],
        notes="Mostrar monitoreo en tiempo real del pipeline.",
    )

    _image_slide(
        prs, "Vista autoridad ambiental",
        CAPTURAS / "dashboard_presentacion.png",
        ["4 hipótesis", "Hallazgos dinámicos", "Contexto por contaminante"],
        notes="Vincular hipótesis con datos de correlación en vivo.",
    )

    _image_slide(
        prs, "Docker y despliegue",
        CAPTURAS / "docker_compose.png",
        [
            "4 servicios orquestados",
            "Healthchecks encadenados",
            "deploy.sh / deploy.ps1",
            ".env para configuración",
        ],
    )

    _content_slide(prs, "Colaboración Git", [
        "Git Flow: main + develop + feature/*",
        "Pull Requests con plantilla y checklist",
        "Convención de commits semánticos",
        "Integrantes: @javier-tc, @roz-ctrl",
        "Evidencia: repo/COLABORACION.md, GitHub PRs",
    ], notes="Mostrar ramas y PR en GitHub si hay conexión.")

    _content_slide(prs, "Testing y automatización", [
        "25 tests pytest — 100% exitosos",
        "72% cobertura etl/ + api/",
        "Reporte HTML: reports/coverage/",
        "Tests: extract, transform, validation, load, API",
        "Mejora futura: CI/CD con GitHub Actions",
    ])

    _content_slide(prs, "Lecciones aprendidas", [
        "Docker Compose acelera demos reproducibles en evaluaciones.",
        "API como capa intermedia simplifica testing del dashboard.",
        "Validación dual detecta problemas antes de la carga.",
        "Git Flow documentado mejora coordinación del equipo.",
        "Pendiente: scheduler ETL, tests E2E, más fuentes en tiempo real.",
    ])

    _content_slide(prs, "Cierre y próximos pasos", [
        "Repositorio: " + REPO_URL,
        "API: http://localhost:8000/docs",
        "Dashboard: http://localhost:8050",
        "Informe ejecutivo: docs/entregables/informe_evaluacion.docx",
        "¿Preguntas?",
    ])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"Presentación generada: {OUT}")


if __name__ == "__main__":
    build()
